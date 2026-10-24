"""ورود اطلاعات افتتاحیه (R314): قالب اکسل/CSV، بارگذاری با پیش‌نمایش خطا، ورود دستی، اعتبارسنجی دوباره.

هر بارگذاری یک نسخه (import_batches) با نام فایل، اثر SHA-256، کاربر و زمان است. ردیف‌ها با کلید طبیعی
هر ماژول ذخیره می‌شوند، پس بارگذاری دوباره ردیف تکراری نمی‌سازد و ردیف ثبت‌قطعی‌شده هرگز بازنویسی نمی‌شود.
هیچ سند مالی یا انباری این‌جا ساخته نمی‌شود.
"""

from __future__ import annotations

import csv
import hashlib
import io
import os
from types import SimpleNamespace

from sqlalchemy import func, select

from peecha.db.base import new_session
from peecha.db.models.opening import ImportBatch, StagedRow
from peecha.services.opening import common as c
from peecha.services.opening import runs
from peecha.services.opening.modules import Lookups, handler, normalize

MAX_ROWS = 50000


def _norm_header(value) -> str:
    return c.text(value).replace("*", "").replace("‌", " ").replace("ـ", "").strip().lower()


def _header_map(h) -> dict[str, str]:
    out = {}
    for col in h.columns:
        out[_norm_header(col.key)] = col.key
        out[_norm_header(col.label)] = col.key
    return out


# --- قالب ------------------------------------------------------------------------------------------------
def write_template(module_code: str, path: str) -> str:
    """قالب خالی: سرستون‌ها (الزامی با ستاره) و برگهٔ راهنما. پسوند csv ← فایل CSV با BOM."""
    h = handler(module_code)
    headers = [col.label + (" *" if col.required else "") for col in h.columns]
    if path.lower().endswith(".csv"):
        with open(path, "w", encoding="utf-8-sig", newline="") as fh:
            csv.writer(fh).writerow(headers)
        return path
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = c.MODULE_LABELS[module_code][:28]
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="FDE9D9")
    for idx, col in enumerate(h.columns, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = bold
        if col.required:
            cell.fill = fill
        ws.column_dimensions[cell.column_letter].width = max(14, len(col.label) + 4)
    guide = wb.create_sheet("راهنما")
    guide.sheet_view.rightToLeft = True
    guide.append(["ستون", "الزامی", "نوع", "توضیح"])
    kinds = {"text": "متن", "decimal": "عدد", "date": "تاریخ شمسی (۱۴۰۳/۰۱/۱۵)", "choice": "یکی از مقادیر مجاز"}
    for col in h.columns:
        hint = col.help
        if col.choices:
            allowed = "، ".join(k for k in col.choices if not k.isascii())
            hint = f"{hint} — مجاز: {allowed}" if hint else f"مجاز: {allowed}"
        guide.append([col.label, "بله" if col.required else "", kinds[col.kind], hint])
    for cell in guide[1]:
        cell.font = bold
    for letter, width in zip("ABCD", (26, 8, 24, 70)):
        guide.column_dimensions[letter].width = width
    wb.save(path)
    return path


def read_file(module_code: str, path: str) -> tuple[list[tuple[int, dict]], str, list[str]]:
    """(ردیف‌ها با شمارهٔ ردیف فایل، اثر SHA-256، خطاهای سرستون). ردیف کاملاً خالی نادیده گرفته می‌شود."""
    h = handler(module_code)
    with open(path, "rb") as fh:
        data = fh.read()
    digest = hashlib.sha256(data).hexdigest()
    if path.lower().endswith(".csv"):
        raw = data.decode("utf-8-sig", errors="replace")
        grid = list(csv.reader(io.StringIO(raw)))
    elif path.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        grid = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
    else:
        raise ValueError("فقط فایل اکسل (xlsx) یا CSV پذیرفته می‌شود.")
    if not grid:
        return [], digest, ["فایل خالی است."]
    names = _header_map(h)
    index: dict[int, str] = {}
    unknown = []
    for pos, cell in enumerate(grid[0]):
        key = names.get(_norm_header(cell))
        if key:
            index[pos] = key
        elif c.text(cell):
            unknown.append(c.text(cell))
    problems = []
    missing = [col.label for col in h.columns if col.required and col.key not in index.values()]
    if missing:
        problems.append("ستون الزامی در فایل نیست: " + "، ".join(missing))
    if unknown:
        problems.append("ستون ناشناخته (نادیده گرفته شد): " + "، ".join(unknown))
    rows = []
    for no, line in enumerate(grid[1:], start=2):
        values = {key: line[pos] for pos, key in index.items() if pos < len(line)}
        if all(c.text(v) == "" for v in values.values()):
            continue
        rows.append((no, values))
    if len(rows) > MAX_ROWS:
        problems.append(f"بیش از {MAX_ROWS:,} ردیف؛ فایل را به چند بخش تقسیم کنید.")
    return rows, digest, problems


# --- اعتبارسنجی ردیف -------------------------------------------------------------------------------------
def _labels(h) -> dict[str, str]:
    return {col.key: col.label for col in h.columns}


def _check(h, lk: Lookups, raw: dict) -> SimpleNamespace:
    payload, values, fmt_errors = normalize(h, raw)
    key_values = {k: (v if v not in (None, "") else payload.get(k, "")) for k, v in values.items()}
    key = h.key(key_values, lk)
    labels = _labels(h)
    if fmt_errors:
        chk_errors, chk_warnings, resolved, amount = fmt_errors, [], None, c.ZERO
    else:
        chk = h.validate(lk, values)
        chk_errors, chk_warnings, resolved, amount = chk.errors, chk.warnings, chk.resolved, chk.amount
        if h.code in _ONCE_ONLY and lk.run.kind_code == "CORRECTION" and _transferred_before(lk, h.code, key):
            chk_errors = chk_errors + [(h.columns[0].key, "این مورد در افتتاحیهٔ قبلی منتقل شده است؛ برای حذف/تغییرش "
                                                          "اول آن را برگشت بزنید")]
    as_list = lambda items: [{"col": col, "label": labels.get(col, col), "msg": msg} for col, msg in items]  # noqa: E731
    return SimpleNamespace(key=key[:300], payload=payload, resolved=c.jsonable(resolved) if resolved else None,
                           errors=as_list(chk_errors), warnings=as_list(chk_warnings), amount=c.money(amount))


_ONCE_ONLY = ("RECEIVABLES", "PAYABLES", "WIP")


def _transferred_before(lk: Lookups, module_code: str, key: str) -> bool:
    from peecha.db.models.opening import OpeningRun

    return bool(lk.s.scalar(select(func.count()).select_from(StagedRow).join(OpeningRun, OpeningRun.run_id == StagedRow.run_id).where(
        OpeningRun.company_id == lk.company_id, StagedRow.run_id != lk.run.run_id, StagedRow.module_code == module_code,
        StagedRow.row_key == key, StagedRow.status_code == "POSTED")))


def _store(row: StagedRow, res: SimpleNamespace, user_id: int, batch_id: int | None, row_no: int | None) -> None:
    row.payload, row.resolved = res.payload, res.resolved
    row.errors, row.warnings = (res.errors or None), (res.warnings or None)
    row.status_code = "INVALID" if res.errors else "VALID"
    row.amount = res.amount
    row.batch_id, row.source_row_no = batch_id, row_no
    row.updated_by_user_id, row.updated_at = user_id, c.now()


def _guard(session, company_id: int, user_id: int, run_id: int, module_code: str):
    run = runs._run(session, company_id, run_id, lock=True)
    if run.status_code not in runs.ENTRY_RUN:
        raise ValueError("ورود اطلاعات فقط پیش از ثبت قطعی کامل فرایند ممکن است.")
    m = runs._module(session, run_id, module_code, lock=True)
    if m.status_code in runs.LOCKED_MODULE:
        raise ValueError("این ماژول ثبت قطعی شده است؛ تغییر فقط با افتتاحیهٔ اصلاحی.")
    if m.owner_user_id and m.owner_user_id != user_id and not c.is_super_admin(user_id):
        raise ValueError("فقط مسئول ورود اطلاعات همین ماژول می‌تواند اطلاعات وارد کند.")
    return run, m


# --- بارگذاری فایل --------------------------------------------------------------------------------------
def import_file(company_id: int, user_id: int, run_id: int, module_code: str, path: str,
                replace_missing: bool = False) -> SimpleNamespace:
    """بارگذاری نسخهٔ تازه. ردیف هم‌کلید به‌روز می‌شود، ردیف ثبت‌شده دست نمی‌خورد؛ replace_missing ردیف‌های
    ثبت‌نشده‌ای را که در فایل تازه نیستند حذف می‌کند."""
    c.require(company_id, user_id, c.FORM_ENTER)
    h = handler(module_code)
    rows, digest, problems = read_file(module_code, path)
    blocking = [p for p in problems if not p.startswith("ستون ناشناخته")]
    if blocking:
        raise ValueError("\n".join(blocking))
    if not rows:
        raise ValueError("فایل هیچ ردیف داده‌ای ندارد.")
    with new_session() as session:
        run, _m = _guard(session, company_id, user_id, run_id, module_code)
        same = session.scalar(select(ImportBatch).where(ImportBatch.run_id == run_id, ImportBatch.module_code == module_code,
                                                        ImportBatch.file_hash == digest))
        if same is not None:
            raise ValueError(f"همین فایل قبلاً بارگذاری شده است (نسخهٔ {same.version_no}).")
        version = (session.scalar(select(func.max(ImportBatch.version_no)).where(
            ImportBatch.run_id == run_id, ImportBatch.module_code == module_code)) or 0) + 1
        batch = ImportBatch(run_id=run_id, module_code=module_code, version_no=version, source_code="FILE",
                            file_name=os.path.basename(path)[:255], file_hash=digest, uploaded_by_user_id=user_id)
        session.add(batch)
        session.flush()
        lk = Lookups(session, company_id, run)
        checked: dict[str, tuple[int, SimpleNamespace]] = {}
        duplicates: list[tuple[int, int]] = []
        for no, raw in rows:
            res = _check(h, lk, raw)
            if res.key in checked:
                first_no, first = checked[res.key]
                duplicates.append((first_no, no))
                first.errors.append({"col": h.columns[0].key, "label": h.columns[0].label,
                                     "msg": f"ردیف {no} فایل همین کلید را دارد؛ ردیف‌های تکراری را یکی کنید"})
                continue
            checked[res.key] = (no, res)
        existing = {r.row_key: r for r in session.scalars(select(StagedRow).where(
            StagedRow.run_id == run_id, StagedRow.module_code == module_code))}
        skipped_posted, valid, invalid = [], 0, 0
        for key, (no, res) in checked.items():
            row = existing.get(key)
            if row is not None and row.status_code in ("POSTED", "REVERSED"):
                skipped_posted.append(no)
                continue
            if row is None:
                row = StagedRow(run_id=run_id, module_code=module_code, row_key=key)
                session.add(row)
            _store(row, res, user_id, batch.batch_id, no)
            valid, invalid = valid + (not res.errors), invalid + bool(res.errors)
        removed = 0
        if replace_missing:
            for key, row in existing.items():
                if key not in checked and row.status_code not in ("POSTED", "REVERSED"):
                    session.delete(row)
                    removed += 1
        batch.row_count, batch.valid_count, batch.invalid_count = len(rows), valid, invalid
        batch.posted_skipped = len(skipped_posted)
        session.flush()
        status = runs.refresh_module_status(session, run_id, module_code, user_id)
        details = {"version": version, "file": batch.file_name, "rows": len(rows), "valid": valid, "invalid": invalid,
                   "duplicates": len(duplicates), "posted_skipped": len(skipped_posted), "removed": removed}
        c.event(session, run_id, "IMPORT", user_id, module_code, batch.file_name, details)
        c.audit(session, company_id, user_id, run_id, "IMPORT", {"module": module_code, "hash": digest, **details})
        session.commit()
        return SimpleNamespace(batch_id=batch.batch_id, version_no=version, row_count=len(rows), valid=valid, invalid=invalid,
                               duplicates=duplicates, posted_skipped=skipped_posted, removed=removed, module_status=status,
                               notes=problems)


# --- ورود دستی ------------------------------------------------------------------------------------------
def save_row(company_id: int, user_id: int, run_id: int, module_code: str, values: dict,
             row_id: int | None = None) -> SimpleNamespace:
    """ثبت یا ویرایش یک ردیف از فرم؛ حتی با خطا ذخیره می‌شود تا کاربر بعداً اصلاح کند."""
    c.require(company_id, user_id, c.FORM_ENTER)
    h = handler(module_code)
    with new_session() as session:
        run, _m = _guard(session, company_id, user_id, run_id, module_code)
        res = _check(h, Lookups(session, company_id, run), values)
        clash = session.scalar(select(StagedRow).where(StagedRow.run_id == run_id, StagedRow.module_code == module_code,
                                                       StagedRow.row_key == res.key))
        if row_id is None:
            if clash is not None:
                raise ValueError("ردیفی با همین مشخصات (کلید) قبلاً وارد شده است؛ همان را ویرایش کنید.")
            row = StagedRow(run_id=run_id, module_code=module_code, row_key=res.key)
            session.add(row)
        else:
            row = session.get(StagedRow, row_id)
            if row is None or row.run_id != run_id or row.module_code != module_code:
                raise ValueError("ردیف نامعتبر است.")
            if row.status_code in ("POSTED", "REVERSED"):
                raise ValueError("ردیف ثبت‌شده ویرایش نمی‌شود.")
            if clash is not None and clash.row_id != row.row_id:
                raise ValueError("ردیف دیگری با همین مشخصات (کلید) وجود دارد.")
            row.row_key = res.key
        _store(row, res, user_id, None, None)
        session.flush()
        runs.refresh_module_status(session, run_id, module_code, user_id)
        c.event(session, run_id, "FORM", user_id, module_code, res.key, {"row_id": row.row_id})
        session.commit()
        return SimpleNamespace(row_id=row.row_id, status=row.status_code, errors=res.errors, warnings=res.warnings,
                               amount=res.amount)


def save_rows(company_id: int, user_id: int, run_id: int, module_code: str, items: list[tuple[int | None, dict]],
              delete_ids=(), batch_id: int | None = None) -> SimpleNamespace:
    """چند ردیف فرم در یک تراکنش (فرم سند افتتاحیه): items = [(شناسهٔ ردیف یا None، مقادیر)]. ردیف بی‌شناسه با همان
    کلید، همان ردیف موجود را به‌روز می‌کند؛ ردیف ثبت‌شده دست نمی‌خورد. ردیف با خطا هم ذخیره می‌شود."""
    c.require(company_id, user_id, c.FORM_ENTER)
    h = handler(module_code)
    with new_session() as session:
        run, _m = _guard(session, company_id, user_id, run_id, module_code)
        lk = Lookups(session, company_id, run)
        by_key = {r.row_key: r for r in session.scalars(select(StagedRow).where(
            StagedRow.run_id == run_id, StagedRow.module_code == module_code))}
        by_id = {r.row_id: r for r in by_key.values()}
        deleted = 0
        for rid in delete_ids:
            row = by_id.pop(rid, None)
            if row is None:
                continue
            if row.status_code in ("POSTED", "REVERSED"):
                raise ValueError("ردیف ثبت‌شده حذف نمی‌شود.")
            by_key.pop(row.row_key, None)
            session.delete(row)
            deleted += 1
        session.flush()
        saved, seen = [], set()
        for row_id, values in items:
            res = _check(h, lk, values)
            if res.key in seen:
                raise ValueError(f"دو ردیف با یک مشخصات وارد شده است ({res.key}).")
            seen.add(res.key)
            clash = by_key.get(res.key)
            row = clash if row_id is None else by_id.get(row_id)
            if row_id is not None and row is None:
                raise ValueError("ردیف نامعتبر است.")
            if row is not None and row.status_code in ("POSTED", "REVERSED"):
                raise ValueError(f"ردیف ثبت‌شده ویرایش نمی‌شود ({res.key}).")
            if clash is not None and row is not None and clash.row_id != row.row_id:
                raise ValueError(f"دو ردیف با یک مشخصات وارد شده است ({res.key}).")
            if row is None:
                row = StagedRow(run_id=run_id, module_code=module_code, row_key=res.key)
                session.add(row)
            elif row.row_key != res.key:
                by_key.pop(row.row_key, None)
                row.row_key = res.key
            by_key[res.key] = row
            _store(row, res, user_id, batch_id, None)
            saved.append((row, res))
        session.flush()
        status = runs.refresh_module_status(session, run_id, module_code, user_id)
        c.event(session, run_id, "FORM", user_id, module_code, f"{len(saved)} ردیف", {"saved": len(saved), "deleted": deleted})
        session.commit()
        out = [SimpleNamespace(row_id=row.row_id, key=res.key, status=row.status_code, errors=res.errors, warnings=res.warnings,
                               amount=res.amount) for row, res in saved]
        return SimpleNamespace(saved=out, deleted=deleted, invalid=sum(1 for x in out if x.errors), module_status=status)


def delete_row(company_id: int, user_id: int, run_id: int, module_code: str, row_id: int) -> None:
    c.require(company_id, user_id, c.FORM_ENTER)
    with new_session() as session:
        _guard(session, company_id, user_id, run_id, module_code)
        row = session.get(StagedRow, row_id)
        if row is None or row.run_id != run_id or row.module_code != module_code:
            raise ValueError("ردیف نامعتبر است.")
        if row.status_code in ("POSTED", "REVERSED"):
            raise ValueError("ردیف ثبت‌شده حذف نمی‌شود.")
        key = row.row_key
        session.delete(row)
        session.flush()
        runs.refresh_module_status(session, run_id, module_code, user_id)
        c.event(session, run_id, "DELETE_ROW", user_id, module_code, key)
        session.commit()


def revalidate(company_id: int, user_id: int, run_id: int, module_code: str) -> SimpleNamespace:
    """اعتبارسنجی دوبارهٔ ردیف‌های ثبت‌نشده (مثلاً پس از تعریف کالا یا حساب جاافتاده)."""
    c.require(company_id, user_id, c.FORM_ENTER)
    h = handler(module_code)
    with new_session() as session:
        run, _m = _guard(session, company_id, user_id, run_id, module_code)
        lk = Lookups(session, company_id, run)
        valid = invalid = 0
        for row in session.scalars(select(StagedRow).where(
                StagedRow.run_id == run_id, StagedRow.module_code == module_code,
                StagedRow.status_code.in_(("VALID", "INVALID", "FAILED"))).order_by(StagedRow.row_id)):
            res = _check(h, lk, row.payload)
            if res.key != row.row_key:
                res.key = row.row_key
            _store(row, res, user_id, row.batch_id, row.source_row_no)
            valid, invalid = valid + (not res.errors), invalid + bool(res.errors)
        session.flush()
        status = runs.refresh_module_status(session, run_id, module_code, user_id)
        c.event(session, run_id, "REVALIDATE", user_id, module_code, None, {"valid": valid, "invalid": invalid})
        session.commit()
        return SimpleNamespace(valid=valid, invalid=invalid, module_status=status)


# --- فهرست و خروجی خطا ------------------------------------------------------------------------------------
def list_rows(company_id: int, run_id: int, module_code: str, status: str | None = None,
              search: str | None = None) -> list[SimpleNamespace]:
    h = handler(module_code)
    with new_session() as session:
        runs._run(session, company_id, run_id)
        q = select(StagedRow).where(StagedRow.run_id == run_id, StagedRow.module_code == module_code)
        if status:
            q = q.where(StagedRow.status_code == status)
        out = []
        needle = c.text(search).lower() if search else ""
        for r in session.scalars(q.order_by(StagedRow.source_row_no.nulls_last(), StagedRow.row_id)):
            if needle and needle not in " ".join(str(v) for v in (r.payload or {}).values()).lower():
                continue
            out.append(SimpleNamespace(
                row_id=r.row_id, row_key=r.row_key, row_no=r.source_row_no, batch_id=r.batch_id, status=r.status_code,
                status_label=c.ROW_STATUS_LABELS.get(r.status_code, r.status_code), payload=dict(r.payload or {}),
                resolved=dict(r.resolved or {}), errors=list(r.errors or []), warnings=list(r.warnings or []),
                amount=c.dec(r.amount), target_type=r.target_type, target_id=r.target_id,
                journal_entry_id=r.journal_entry_id, posted_at=r.posted_at,
                problem_text="؛ ".join(f"{e['label']}: {e['msg']}" for e in (r.errors or [])),
                warning_text="؛ ".join(f"{w['label']}: {w['msg']}" for w in (r.warnings or []))))
        _ = h
        return out


def export_rows(company_id: int, run_id: int, module_code: str, path: str, only_errors: bool = False) -> str:
    """همان قالب ورود + ستون «خطا/هشدار»؛ پس از اصلاح، همین فایل دوباره بارگذاری می‌شود."""
    import openpyxl

    h = handler(module_code)
    rows = list_rows(company_id, run_id, module_code, "INVALID" if only_errors else None)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = c.MODULE_LABELS[module_code][:28]
    ws.sheet_view.rightToLeft = True
    ws.append([col.label + (" *" if col.required else "") for col in h.columns] + ["وضعیت", "خطا", "هشدار"])
    for r in rows:
        ws.append([r.payload.get(col.key, "") for col in h.columns] + [r.status_label, r.problem_text, r.warning_text])
    wb.save(path)
    return path


def list_batches(company_id: int, run_id: int, module_code: str | None = None) -> list[SimpleNamespace]:
    with new_session() as session:
        runs._run(session, company_id, run_id)
        q = select(ImportBatch).where(ImportBatch.run_id == run_id)
        if module_code:
            q = q.where(ImportBatch.module_code == module_code)
        items = session.scalars(q.order_by(ImportBatch.module_code, ImportBatch.version_no.desc())).all()
        names = c.user_names(session, [b.uploaded_by_user_id for b in items])
        return [SimpleNamespace(batch_id=b.batch_id, module_code=b.module_code, module_label=c.MODULE_LABELS[b.module_code],
                                version_no=b.version_no, file_name=b.file_name or "", file_hash=b.file_hash or "",
                                rows=b.row_count, valid=b.valid_count, invalid=b.invalid_count, posted_skipped=b.posted_skipped,
                                user=names.get(b.uploaded_by_user_id, ""), at=b.uploaded_at) for b in items]
