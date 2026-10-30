"""R375: گزارش زمان‌بندی‌شده.

- هر گزارش ذخیره‌شده می‌تواند چند زمان‌بندی روزانه، هفتگی یا ماهانه (ساعت مشخص) داشته باشد، با بازهٔ تاریخ نسبی روی ستون
  تاریخ داده (دیروز، هفت روز گذشته، ماه جاری یا ماه گذشتهٔ شمسی) و گیرندگانی از کاربران همین شرکت.
- اجرا برای هر گیرنده با دسترسی خود او است: گیرنده‌ای که به آن داده دسترسی ندارد کنار گذاشته می‌شود و گیرندگانی که شعبه‌های
  مجاز متفاوت دارند فایل جدا می‌گیرند. فایل اکسل در خود پایگاه داده ذخیره و از مرکز اعلان خبر داده می‌شود؛ خطا به سازنده.
- run_due وقتی برنامهٔ دسکتاپ باز است هر چند دقیقه صدا زده می‌شود (یا با زمان‌بند سیستم‌عامل:
  python -m peecha.services.report_builder.schedule). هر زمان‌بندی پیش از اجرا با به‌روزرسانی شرطی «برداشته» می‌شود تا اگر
  چند رایانه هم‌زمان باز بودند فقط یکی آن را اجرا کند.
"""

from __future__ import annotations

import datetime
import io
from dataclasses import dataclass, field

import jdatetime
from sqlalchemy import select, update

from peecha.db.base import new_session
from peecha.db.models.reporting import CustomReport, ReportRun, ReportSchedule
from peecha.services.report_builder import engine
from peecha.services.report_builder.datasets import DATASETS

FREQUENCIES = {"DAILY": "روزانه", "WEEKLY": "هفتگی", "MONTHLY": "ماهانه"}
PERIODS = {"NONE": "بدون بازهٔ تاریخ", "YESTERDAY": "دیروز", "LAST_7_DAYS": "هفت روز گذشته", "THIS_MONTH": "ماه جاری",
           "LAST_MONTH": "ماه گذشته"}
# روزهای هفته به ترتیب ایرانی؛ مقدار همان weekday پایتون (دوشنبه = ۰)
WEEKDAYS = [(5, "شنبه"), (6, "یکشنبه"), (0, "دوشنبه"), (1, "سه‌شنبه"), (2, "چهارشنبه"), (3, "پنجشنبه"), (4, "جمعه")]
NOTIFY_TYPE = "RPT_SCHEDULED"
MAX_FILE_BYTES = 15 * 1024 * 1024


@dataclass
class ScheduleFields:
    frequency_code: str
    run_hour: int
    recipient_user_ids: list[int]
    weekday: int | None = None
    month_day: int | None = None
    period_code: str = "NONE"
    is_active: bool = True


@dataclass
class ScheduleRow:
    schedule_id: int
    report_id: int
    fields: ScheduleFields
    next_run_at: datetime.datetime
    last_run_at: datetime.datetime | None

    @property
    def label(self) -> str:
        f = self.fields
        when = FREQUENCIES[f.frequency_code]
        if f.frequency_code == "WEEKLY":
            when += f" ({dict(WEEKDAYS).get(f.weekday, '')})"
        elif f.frequency_code == "MONTHLY":
            from peecha import numerals

            when += f" (روز {numerals.to_persian_digits(str(f.month_day))})"
        return when


@dataclass
class RunRow:
    run_id: int
    report_id: int
    report_name: str
    schedule_id: int | None
    run_at: datetime.datetime
    status_code: str
    row_count: int
    file_name: str
    error_message: str
    recipient_user_ids: list[int] = field(default_factory=list)


def _local_now() -> datetime.datetime:
    return datetime.datetime.now().astimezone()


def next_run(f: ScheduleFields, after: datetime.datetime) -> datetime.datetime:
    """اولین زمان اجرای بعد از after (به وقت محلی)."""
    after = after.astimezone()
    day = after.date()
    for _ in range(400):
        due = after.replace(year=day.year, month=day.month, day=day.day, hour=f.run_hour, minute=0, second=0, microsecond=0)
        ok = (f.frequency_code == "DAILY" or (f.frequency_code == "WEEKLY" and day.weekday() == f.weekday)
              or (f.frequency_code == "MONTHLY" and jdatetime.date.fromgregorian(date=day).day == f.month_day))
        if ok and due > after:
            return due
        day += datetime.timedelta(days=1)
    raise ValueError("زمان اجرای بعدی پیدا نشد.")


def period_range(code: str, on: datetime.date) -> tuple[datetime.date, datetime.date] | None:
    """بازهٔ تاریخ نسبی به روز اجرا؛ ماه‌ها شمسی."""
    if code == "NONE":
        return None
    if code == "YESTERDAY":
        d = on - datetime.timedelta(days=1)
        return d, d
    if code == "LAST_7_DAYS":
        return on - datetime.timedelta(days=7), on - datetime.timedelta(days=1)
    j = jdatetime.date.fromgregorian(date=on)
    first = jdatetime.date(j.year, j.month, 1)
    if code == "THIS_MONTH":
        return first.togregorian(), on
    last_end = first.togregorian() - datetime.timedelta(days=1)
    lj = jdatetime.date.fromgregorian(date=last_end)
    return jdatetime.date(lj.year, lj.month, 1).togregorian(), last_end


def _validate(company_id: int, f: ScheduleFields, report: CustomReport) -> list[int]:
    from peecha.db.models.security import UserCompany

    if f.frequency_code not in FREQUENCIES or f.period_code not in PERIODS:
        raise ValueError("تکرار یا بازهٔ تاریخ نامعتبر است.")
    if not 0 <= int(f.run_hour) <= 23:
        raise ValueError("ساعت اجرا از ۰ تا ۲۳ است.")
    if f.frequency_code == "WEEKLY" and f.weekday not in range(7):
        raise ValueError("روز هفته را انتخاب کنید.")
    if f.frequency_code == "MONTHLY" and not (f.month_day and 1 <= int(f.month_day) <= 28):
        raise ValueError("روز ماه از ۱ تا ۲۸ است.")
    if f.period_code != "NONE" and DATASETS[report.dataset_code].date_key is None:
        raise ValueError("این داده ستون تاریخ ندارد؛ بازهٔ تاریخ را «بدون بازه» بگذارید.")
    recipients = sorted({int(u) for u in f.recipient_user_ids if u})
    if not recipients:
        raise ValueError("دست‌کم یک گیرنده انتخاب کنید.")
    with new_session() as session:
        members = set(session.scalars(select(UserCompany.user_id).where(UserCompany.company_id == company_id,
                                                                        UserCompany.user_id.in_(recipients))))
    outsiders = [u for u in recipients if u not in members]
    if outsiders:
        raise ValueError("گیرنده باید از کاربران همین شرکت باشد.")
    without = [u for u in recipients if not engine.can_use(company_id, u, report.dataset_code)]
    if without:
        raise ValueError("بعضی گیرندگان به داده‌های این گزارش دسترسی ندارند.")
    return recipients


def save_schedule(company_id: int, user_id: int, report_id: int, f: ScheduleFields, schedule_id: int | None = None,
                  now: datetime.datetime | None = None) -> int:
    saved = engine.get_report(company_id, user_id, report_id)
    if not saved.editable_by(company_id, user_id):
        raise ValueError("فقط سازندهٔ گزارش یا مدیر سیستم آن را زمان‌بندی می‌کند.")
    with new_session() as session:
        report = session.get(CustomReport, report_id)
        recipients = _validate(company_id, f, report)
        if schedule_id is None:
            row = ReportSchedule(company_id=company_id, report_id=report_id, created_by_user_id=user_id)
            session.add(row)
        else:
            row = session.get(ReportSchedule, schedule_id)
            if row is None or row.company_id != company_id or row.report_id != report_id:
                raise ValueError("زمان‌بندی نامعتبر است.")
        row.frequency_code, row.run_hour, row.period_code = f.frequency_code, int(f.run_hour), f.period_code
        row.weekday = f.weekday if f.frequency_code == "WEEKLY" else None
        row.month_day = f.month_day if f.frequency_code == "MONTHLY" else None
        row.recipient_user_ids, row.is_active = recipients, bool(f.is_active)
        row.next_run_at = next_run(f, now or _local_now())
        session.commit()
        return row.schedule_id


def delete_schedule(company_id: int, user_id: int, schedule_id: int) -> None:
    with new_session() as session:
        row = session.get(ReportSchedule, schedule_id)
        if row is None or row.company_id != company_id:
            raise ValueError("زمان‌بندی نامعتبر است.")
        report_id = row.report_id
    if not engine.get_report(company_id, user_id, report_id).editable_by(company_id, user_id):
        raise ValueError("فقط سازندهٔ گزارش یا مدیر سیستم زمان‌بندی را حذف می‌کند.")
    with new_session() as session:
        session.delete(session.get(ReportSchedule, schedule_id))
        session.commit()


def list_schedules(company_id: int, report_id: int) -> list[ScheduleRow]:
    with new_session() as session:
        rows = list(session.scalars(select(ReportSchedule).where(ReportSchedule.company_id == company_id,
                                                                 ReportSchedule.report_id == report_id)
                                    .order_by(ReportSchedule.schedule_id)))
    return [ScheduleRow(r.schedule_id, r.report_id, ScheduleFields(r.frequency_code, r.run_hour, list(r.recipient_user_ids or []),
                                                                   r.weekday, r.month_day, r.period_code, r.is_active),
                        r.next_run_at, r.last_run_at) for r in rows]


# --- خروجی اکسل -----------------------------------------------------------------------------------------------
def to_xlsx(result: engine.Result, company_name: str = "", run_at: datetime.datetime | None = None) -> bytes:
    import openpyxl
    from openpyxl.styles import Font

    from peecha import numerals

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = (result.title[:31] or "گزارش").replace("/", "-")
    sheet.sheet_view.rightToLeft = True
    sheet.append([result.title])
    sheet.cell(row=1, column=1).font = Font(bold=True, size=13)
    when = run_at or _local_now()
    sheet.append([f"{company_name} — {numerals.format_jalali_datetime(when)}".strip(" —")])
    for label, text in result.filters_text:
        sheet.append([f"{label}: {text}"])
    sheet.append([f"منبع: {result.source}"])
    sheet.append([])
    header_row = sheet.max_row + 1
    sheet.append(result.headers)
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)

    def cell(v):
        return numerals.format_jalali_date(v) if isinstance(v, datetime.date) else v

    for line in result.rows:
        sheet.append([cell(v) for v in line])
    if result.totals:
        sheet.append(["جمع" if i == 0 and v is None else cell(v) for i, v in enumerate(result.totals)])
        for c in sheet[sheet.max_row]:
            c.font = Font(bold=True)
    sheet.freeze_panes = f"A{header_row + 1}"
    for col_cells in sheet.iter_cols(min_row=header_row):
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        sheet.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 45)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- اجرا ------------------------------------------------------------------------------------------------------
def _definition_for(saved: engine.SavedReport, period_code: str, on: datetime.date) -> engine.Definition:
    definition = engine.Definition.from_json(saved.definition.to_json())
    span = period_range(period_code, on)
    date_key = DATASETS[saved.dataset_code].date_key
    if span and date_key:
        definition.filters.append(engine.Filter(date_key, "BETWEEN", span[0], span[1]))
    return definition


def _company_name(company_id: int) -> str:
    from peecha.db.models.core import Company

    with new_session() as session:
        company = session.get(Company, company_id)
        return (company.display_name or company.legal_name or "") if company else ""


def _groups(company_id: int, dataset_code: str, recipients: list[int]) -> dict[frozenset | None, list[int]]:
    """گیرندگان مجاز، دسته‌بندی‌شده با شعبه‌های مجاز (هر دسته یک فایل)."""
    from peecha.services import branch_access

    out: dict[frozenset | None, list[int]] = {}
    for uid in recipients:
        if engine.can_use(company_id, uid, dataset_code):
            out.setdefault(branch_access.allowed_branch_ids(company_id, uid), []).append(uid)
    return out


def execute(company_id: int, report_id: int, period_code: str, recipients: list[int], *, schedule_id: int | None = None,
            owner_user_id: int | None = None, now: datetime.datetime | None = None) -> list[int]:
    """گزارش را برای گیرندگان اجرا، ذخیره و اعلان می‌کند؛ شناسهٔ اجراها را برمی‌گرداند."""
    from peecha.services.workflow import notify

    now = now or _local_now()
    with new_session() as session:
        report = session.get(CustomReport, report_id)
        if report is None or report.company_id != company_id:
            raise ValueError("گزارش نامعتبر است.")
        name, dataset_code, owner = report.name, report.dataset_code, owner_user_id or report.owner_user_id
    saved = engine.SavedReport(report_id, name, dataset_code, "", owner, False, engine.Definition.from_json(report.definition), now)
    definition = _definition_for(saved, period_code, now.date())
    run_ids = []
    groups = _groups(company_id, dataset_code, recipients)
    stamp = jdatetime.datetime.fromgregorian(datetime=now).strftime("%Y-%m-%d-%H%M")
    company_name = _company_name(company_id)
    for _branches, users in groups.items():
        try:
            result = engine.run(company_id, users[0], definition, name)
            content = to_xlsx(result, company_name, now)
            if len(content) > MAX_FILE_BYTES:
                raise ValueError("فایل گزارش بیش از اندازه بزرگ است؛ فیلتر یا بازهٔ کوتاه‌تری بگذارید.")
            row = ReportRun(company_id=company_id, schedule_id=schedule_id, report_id=report_id, run_at=now, status_code="OK",
                            row_count=len(result.rows), recipient_user_ids=users, file_name=f"{name}-{stamp}.xlsx", content=content)
            error = None
        except Exception as exc:  # noqa: BLE001 -- خطای یک اجرا ثبت و به سازنده گفته می‌شود؛ زمان‌بندی‌های دیگر ادامه می‌دهند
            error = str(exc)[:500] or exc.__class__.__name__
            row = ReportRun(company_id=company_id, schedule_id=schedule_id, report_id=report_id, run_at=now, status_code="ERROR",
                            recipient_user_ids=users, error_message=error)
        with new_session() as session:
            session.add(row)
            session.commit()
            run_ids.append(row.run_id)
        if error is None:
            notify.send(company_id, users, NOTIFY_TYPE, f"گزارش «{name}» آماده است",
                        f"{row.row_count} ردیف؛ فایل اکسل در گزارش‌ساز، بخش اجراهای زمان‌بندی‌شده.", entity_type="ReportRun",
                        entity_id=row.run_id)
        else:
            notify.send(company_id, [owner], NOTIFY_TYPE, f"اجرای گزارش «{name}» انجام نشد", error, entity_type="ReportRun",
                        entity_id=row.run_id)
    if not groups:
        notify.send(company_id, [owner], NOTIFY_TYPE, f"گزارش «{name}» گیرنده‌ای ندارد",
                    "هیچ‌یک از گیرندگان دیگر به داده‌های این گزارش دسترسی ندارند.", entity_type="CustomReport", entity_id=report_id)
    return run_ids


def run_now(company_id: int, user_id: int, schedule_id: int) -> list[int]:
    with new_session() as session:
        row = session.get(ReportSchedule, schedule_id)
        if row is None or row.company_id != company_id:
            raise ValueError("زمان‌بندی نامعتبر است.")
        report_id, period, recipients = row.report_id, row.period_code, list(row.recipient_user_ids or [])
    if not engine.get_report(company_id, user_id, report_id).editable_by(company_id, user_id):
        raise ValueError("فقط سازندهٔ گزارش یا مدیر سیستم آن را اجرا می‌کند.")
    return execute(company_id, report_id, period, recipients, schedule_id=schedule_id)


def run_due(now: datetime.datetime | None = None) -> int:
    """زمان‌بندی‌های سررسیدهٔ همهٔ شرکت‌ها؛ تعداد اجراشده."""
    now = (now or _local_now()).astimezone()
    with new_session() as session:
        due = list(session.scalars(select(ReportSchedule).where(ReportSchedule.is_active.is_(True), ReportSchedule.next_run_at <= now)
                                   .order_by(ReportSchedule.next_run_at)))
        due = [(r.schedule_id, r.company_id, r.report_id, r.period_code, list(r.recipient_user_ids or []), r.next_run_at,
                ScheduleFields(r.frequency_code, r.run_hour, [], r.weekday, r.month_day, r.period_code)) for r in due]
    count = 0
    for schedule_id, company_id, report_id, period, recipients, old_next, fields in due:
        with new_session() as session:
            claimed = session.execute(update(ReportSchedule).where(ReportSchedule.schedule_id == schedule_id,
                                                                   ReportSchedule.next_run_at == old_next)
                                      .values(next_run_at=next_run(fields, now), last_run_at=now)).rowcount
            session.commit()
        if claimed != 1:
            continue
        execute(company_id, report_id, period, recipients, schedule_id=schedule_id, now=now)
        count += 1
    return count


def list_runs(company_id: int, user_id: int, report_id: int | None = None, limit: int = 100) -> list[RunRow]:
    """اجراهایی که کاربر گیرنده‌شان است (سازندهٔ گزارش و مدیر سیستم همه را می‌بینند)."""
    from peecha.services import roles

    manager = roles.is_manager(user_id, company_id)
    with new_session() as session:
        q = (select(ReportRun.run_id, ReportRun.report_id, CustomReport.name, CustomReport.owner_user_id, ReportRun.schedule_id,
                    ReportRun.run_at, ReportRun.status_code, ReportRun.row_count, ReportRun.file_name, ReportRun.error_message,
                    ReportRun.recipient_user_ids)
             .join(CustomReport, CustomReport.report_id == ReportRun.report_id).where(ReportRun.company_id == company_id))
        if report_id is not None:
            q = q.where(ReportRun.report_id == report_id)
        rows = session.execute(q.order_by(ReportRun.run_at.desc(), ReportRun.run_id.desc()).limit(limit)).all()
    return [RunRow(r[0], r[1], r[2], r[4], r[5], r[6], r[7], r[8] or "", r[9] or "", list(r[10] or [])) for r in rows
            if manager or r[3] == user_id or user_id in (r[10] or [])]


def run_file(company_id: int, user_id: int, run_id: int) -> tuple[str, bytes]:
    if run_id not in {r.run_id for r in list_runs(company_id, user_id, limit=100000)}:
        raise ValueError("این فایل برای شما فرستاده نشده است.")
    with new_session() as session:
        row = session.get(ReportRun, run_id)
        if row.content is None:
            raise ValueError("این اجرا فایلی ندارد.")
        return row.file_name, bytes(row.content)


if __name__ == "__main__":  # زمان‌بند سیستم‌عامل
    from peecha.db.base import get_engine
    from peecha.db.schema_bootstrap import apply_pending_schema_files

    apply_pending_schema_files(get_engine())
    print(run_due())
