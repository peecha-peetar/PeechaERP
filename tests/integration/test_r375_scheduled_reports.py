"""R375: گزارش زمان‌بندی‌شده: زمان اجرای بعدی (روزانه، هفتگی، ماه شمسی)، بازهٔ نسبی، اعتبارسنجی گیرندگان، اجرای سررسیده فقط
یک بار، فایل جدا برای گیرندهٔ محدود به شعبه، اعلان، دیدن فایل فقط برای گیرنده، خطای اجرا برای سازنده، حذف آبشاری؛ و فرم."""
import io
import os, sys
os.environ["PEECHA_DB_NAME"] = "peecha_test_r375"
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from prd_fixture import *  # noqa: F401,F403

import jdatetime
import openpyxl
from sqlalchemy import update
from peecha.db.models.commercial import CommercialDocument
from peecha.db.models.reporting import CustomReport, ReportRun, ReportSchedule
from peecha.db.models.security import Notification
from peecha.services import branch_access, commercial_documents as documents_service, procurement_masters as masters
from peecha.services import roles as roles_service, users as users_service
from peecha.services.report_builder import engine, schedule as rs

LOCAL = lambda *a: datetime.datetime(*a).astimezone()  # noqa: E731
F = rs.ScheduleFields
check(rs.next_run(F("DAILY", 7, []), LOCAL(2026, 10, 30, 8, 0)) == LOCAL(2026, 10, 31, 7, 0), "daily: the next morning")
check(rs.next_run(F("DAILY", 7, []), LOCAL(2026, 10, 30, 6, 0)) == LOCAL(2026, 10, 30, 7, 0), "daily: later today")
check(rs.next_run(F("WEEKLY", 9, [], weekday=5), LOCAL(2026, 10, 30, 8, 0)) == LOCAL(2026, 10, 31, 9, 0), "weekly: next Saturday")
first = jdatetime.date(1405, 9, 1).togregorian()
check(rs.next_run(F("MONTHLY", 6, [], month_day=1), LOCAL(2026, 10, 30, 8, 0)).date() == first, "monthly: first of the next Jalali month")
on = jdatetime.date(1405, 8, 15).togregorian()
check(rs.period_range("YESTERDAY", on) == (on - datetime.timedelta(days=1),) * 2, "yesterday")
check(rs.period_range("THIS_MONTH", on) == (jdatetime.date(1405, 8, 1).togregorian(), on), "this Jalali month to date")
check(rs.period_range("LAST_MONTH", on) == (jdatetime.date(1405, 7, 1).togregorian(), jdatetime.date(1405, 7, 30).togregorian()),
      "the whole previous Jalali month")

b1 = masters.save_branch(company_id, "B1", "شعبهٔ تهران")
b2 = masters.save_branch(company_id, "B2", "شعبهٔ اصفهان")
w1 = locations_service.create_warehouse(company_id, "W1", "فروش تهران", WF(allow_negative_stock=True))
w2 = locations_service.create_warehouse(company_id, "W2", "فروش اصفهان", WF(allow_negative_stock=True))
masters.set_warehouse_branch(company_id, w1, b1)
masters.set_warehouse_branch(company_id, w2, b2)
customer = dimensions_service.create_customer(company_id, "C1", "مشتری")


def invoice(wh, day, qty):
    d = documents_service.create_document(company_id, uid, "SALES_INVOICE", day, documents_service.DocumentHeaderFields(
        counterparty_detail_account_id=customer, currency_id=company.base_currency_id, warehouse_id=wh))
    documents_service.add_line(d, company_id, item_id=fg, uom_id=pcs, quantity=D(qty), quantity_base=D(qty), unit_price=D(1000))
    with new_session() as s:
        s.execute(update(CommercialDocument).where(CommercialDocument.document_id == d).values(status_code="CONFIRMED"))
        s.commit()


invoice(w1, today, 6)
invoice(w2, today, 4)
invoice(w1, today - datetime.timedelta(days=10), 5)

forms = {f.code: f.form_id for f in roles_service.list_forms()}
role = roles_service.create_role(company_id, "فروشنده", None)
roles_service.set_role_permission(role.role_id, forms["commercial_documents_list_sales"], "VIEW", True)


def user(name, sales=True):
    u = users_service.create_user(name, name, "secret123", None, company.default_language_id, False, [company_id], company_id).user_id
    if sales:
        roles_service.set_user_role(u, role.role_id, company_id, True)
    return u


seller, seller1, staff = user("seller"), user("seller1"), user("staff", sales=False)
branch_access.set_user_branches(seller1, company_id, [b1])
rid = engine.save_report(company_id, uid, "فروش روزانه", engine.Definition("sales_lines", ["branch", "quantity", "net_amount"]))
stock_rid = engine.save_report(company_id, uid, "موجودی", engine.Definition("stock", ["item_code", "on_hand"]))
check(raises(lambda: rs.save_schedule(company_id, uid, rid, F("DAILY", 7, [])), "گیرنده"), "recipients required")
check(raises(lambda: rs.save_schedule(company_id, uid, rid, F("DAILY", 7, [uid, staff])), "دسترسی"), "recipients need the data")
check(raises(lambda: rs.save_schedule(company_id, uid, rid, F("WEEKLY", 7, [uid])), "روز هفته"), "weekly needs a weekday")
check(raises(lambda: rs.save_schedule(company_id, uid, stock_rid, F("DAILY", 7, [uid], period_code="YESTERDAY")), "تاریخ"),
      "no date range on data without dates")
check(raises(lambda: rs.save_schedule(company_id, seller, rid, F("DAILY", 7, [seller])), "پیدا نشد"),
      "only the owner schedules (a private report is not even visible)")

base = LOCAL(today.year, today.month, today.day, 8, 0)
sid = rs.save_schedule(company_id, uid, rid, F("DAILY", 7, [uid, seller, seller1], period_code="YESTERDAY"), now=base)
due = rs.list_schedules(company_id, rid)[0].next_run_at
check(due == base.replace(hour=7) + datetime.timedelta(days=1), "next run tomorrow at seven")
check(rs.run_due(now=due - datetime.timedelta(minutes=1)) == 0, "nothing before its time")
check(rs.run_due(now=due) == 1 and rs.run_due(now=due) == 0, "a due schedule runs once")
with new_session() as s:
    runs = list(s.scalars(select(ReportRun).where(ReportRun.schedule_id == sid).order_by(ReportRun.run_id)))
    sched = s.get(ReportSchedule, sid)
    check(sched.next_run_at == due + datetime.timedelta(days=1) and sched.last_run_at == due, "next run moves a day ahead")
check(len(runs) == 2 and all(r.status_code == "OK" for r in runs), "one file per branch scope")
by_users = {tuple(r.recipient_user_ids): r for r in runs}
check(set(by_users) == {(uid, seller), (seller1,)}, f"recipients grouped by branch access ({list(by_users)})")


def sheet_rows(run):
    ws = openpyxl.load_workbook(io.BytesIO(run.content)).active
    return [[c.value for c in row] for row in ws.iter_rows()]


full, limited = sheet_rows(by_users[(uid, seller)]), sheet_rows(by_users[(seller1,)])
check(full[0][0] == "فروش روزانه" and any(str(r[0]).startswith("تاریخ:") for r in full[:5]), "title and the date filter on top")
data_full = [r for r in full if r and isinstance(r[1], (int, float, decimal.Decimal)) and r[0] != "جمع"]
data_limited = [r for r in limited if r and isinstance(r[1], (int, float, decimal.Decimal)) and r[0] != "جمع"]
check(sorted(float(r[1]) for r in data_full) == [4.0, 6.0] and [(r[0], float(r[1])) for r in data_limited] == [("شعبهٔ تهران", 6.0)],
      f"the run day's yesterday only, and the branch-limited file ({data_full}, {data_limited})")
with new_session() as s:
    notes = {u: s.scalar(select(func.count()).select_from(Notification).where(Notification.user_id == u, Notification.type_code == "RPT_SCHEDULED"))
             for u in (uid, seller, seller1, staff)}
check(notes == {uid: 1, seller: 1, seller1: 1, staff: 0}, f"recipients are notified ({notes})")
check([r.run_id for r in rs.list_runs(company_id, seller1)] == [by_users[(seller1,)].run_id] and rs.list_runs(company_id, staff) == [],
      "each user sees only the files sent to them")
check(raises(lambda: rs.run_file(company_id, seller1, by_users[(uid, seller)].run_id), "فرستاده نشده"), "no access to other files")
check(rs.run_file(company_id, seller, by_users[(uid, seller)].run_id)[0].startswith("فروش روزانه-"), "file name")

# خطای اجرا به سازنده
with new_session() as s:
    s.execute(update(CustomReport).where(CustomReport.report_id == rid).values(definition={"dataset": "sales_lines", "columns": ["nope"]}))
    s.commit()
ids = rs.run_now(company_id, uid, sid)
with new_session() as s:
    failed = [s.get(ReportRun, i) for i in ids]
    owner_notes = s.scalar(select(func.count()).select_from(Notification).where(Notification.user_id == uid,
                                                                               Notification.type_code == "RPT_SCHEDULED"))
check(all(r.status_code == "ERROR" and "ستون" in r.error_message for r in failed) and owner_notes == 1 + len(ids),
      "a failing run is recorded and the owner is told")
engine.save_report(company_id, uid, "فروش روزانه", engine.Definition("sales_lines", ["branch", "quantity", "net_amount"]), report_id=rid)

# فرم
from PySide6.QtWidgets import QApplication, QPushButton
qapp = QApplication.instance() or QApplication([])
from peecha.ui.screens.report_builder import ReportBuilderScreen

screen = ReportBuilderScreen()
screen.refresh()
row = next(i for i, r in enumerate(screen.reports) if r.report_id == rid)
screen.saved_table.setCurrentCell(row, 0)
panel = screen.schedule_panel
check(panel.report_id == rid and panel.schedule_table.rowCount() == 1 and panel.run_table.rowCount() == 4, "schedules and runs of the report")
panel.schedule_table.setCurrentCell(0, 0)
check(panel.schedule_id == sid and sorted(panel._recipients()) == sorted([uid, seller, seller1]), "schedule loaded with its recipients")
check(panel.run_now() and panel.run_table.rowCount() == 6, f"run now from the form ({screen.status_label.text()})")
target = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_r375_out.xlsx")
panel.run_table.setCurrentCell(0, 0)
saved_path = panel.save_file(target)
check(saved_path == target and openpyxl.load_workbook(target).active["A1"].value == "فروش روزانه", "file saved from the form")
os.remove(target)
panel.new_schedule()
from peecha.ui.screens.fixed_assets import set_combo
set_combo(panel.frequency_combo, "MONTHLY")
panel.month_day_field.setValue(1)
set_combo(panel.period_combo, "LAST_MONTH")
check(panel.save() and panel.schedule_table.rowCount() == 2, f"monthly schedule from the form ({screen.status_label.text()})")
check(all(len(b.text()) <= 2 for b in screen.findChildren(QPushButton) if b.text()), "icon buttons only")

engine.delete_report(company_id, uid, rid)
with new_session() as s:
    left = (s.scalar(select(func.count()).select_from(ReportSchedule)), s.scalar(select(func.count()).select_from(ReportRun)))
check(left == (0, 0), "deleting the report removes its schedules and files")

finish()
