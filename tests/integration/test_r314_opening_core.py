"""R314 مرحلهٔ A+B: مهاجرت، تنظیمات و حساب واسط، فرایند و وضعیت‌ها، دسترسی و تفکیک وظایف، قالب و بارگذاری
(خطای ستونی، تکرار در فایل، همان فایل دوباره، بارگذاری اصلاحی بدون ردیف تکراری)، ورود دستی، حفاظت ردیف ثبت‌شده،
بارگذاری هم‌زمان و لغو فرایند."""
import os, sys, tempfile, threading
os.environ["PEECHA_DB_NAME"] = "peecha_test_r314_core"
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from opn_fixture import *  # noqa: F401,F403
import opn_fixture as fx
from sqlalchemy.exc import DBAPIError
from peecha.db.models.audit import ActivityLog
from peecha.db.models.opening import ImportBatch, OpeningEvent, RunModule

check, raises = fx.check, fx.raises
tmp = tempfile.mkdtemp(prefix="r314_")

# ===== ۱) مهاجرت =====
with new_session() as s:
    tables = set(s.scalars(text("SELECT table_name FROM information_schema.tables WHERE table_schema = 'opn'")))
    doc_check = s.scalar(text("SELECT pg_get_constraintdef(oid) FROM pg_constraint WHERE conname = 'commercial_documents_document_type_code_check'"))
    txn_check = s.scalar(text("SELECT pg_get_constraintdef(oid) FROM pg_constraint WHERE conname = 'order_transactions_txn_type_check'"))
    feat = s.scalar(text("SELECT count(*) FROM inv.feature_definitions WHERE feature_code = 'OPN_SELF_APPROVAL'"))
check({"settings", "runs", "run_modules", "import_batches", "staged_rows", "postings", "events"} <= tables, f"opn tables exist ({sorted(tables)})")
check(doc_check and "OPENING_RECEIVABLE" in doc_check and "OPENING_PAYABLE" in doc_check, "commercial document types include opening open items")
check(txn_check and "'OPENING'" in txn_check and "'ISSUE'" in txn_check, "production txn types include OPENING (old values kept)")
check(feat == 1, "self-approval toggle registered (off by default)")
check(os.path.exists(os.path.join(os.path.dirname(__file__), "..", "..", "db", "rollback", "224_opening_balances_down.sql")), "rollback script shipped")

# ===== ۲) کاربران و دسترسی =====
entry = user("opn_entry", "OPN_ENTRY")
entry2 = user("opn_entry2", "OPN_ENTRY")
approver = user("opn_appr", "OPN_APPROVER")
manager = user("opn_mgr", "OPN_MANAGER")
viewer = user("opn_view", "OPN_VIEWER")
outsider = user("opn_out")
check(raises(lambda: orun.create_run(company_id, outsider, "X", "x", fy_id, balance_date, ["GL"]), "دسترسی"),
      "user without the setup permission cannot create a run")
check(raises(lambda: orun.create_run(company_id, entry, "X", "x", fy_id, balance_date, ["GL"]), "دسترسی"),
      "data-entry role cannot create a run")
check(raises(lambda: oc.save_settings(company_id, entry, clearing_gl.account_id, "BLOCK", None), "دسترسی"),
      "data-entry role cannot change opening settings")

# ===== ۳) تنظیمات و حساب واسط =====
check(raises(lambda: oc.save_settings(company_id, uid, k_clr.account_id, "BLOCK", None), "معین فعال"),
      "clearing account must be postable")
check(raises(lambda: oc.save_settings(company_id, uid, clearing_gl.account_id, "APPROVAL", None), "اختلاف"),
      "APPROVAL policy needs a difference account")
check(raises(lambda: oc.save_settings(company_id, uid, clearing_gl.account_id, "APPROVAL", clearing_gl.account_id), "همان حساب واسط"),
      "difference account differs from the clearing account")
oc.save_settings(company_id, uid, clearing_gl.account_id, "BLOCK", None)
check(oc.clearing_account_id(company_id) == clearing_gl.account_id and oc.settings(company_id).difference_policy == "BLOCK",
      "clearing account stored in the inventory role mappings; policy BLOCK")
with new_session() as s:
    logged = s.scalar(select(func.count()).select_from(ActivityLog).where(ActivityLog.entity_type == oc.ENTITY_SETTINGS))
check(logged >= 1, "settings change audited")
check("OPENING_CLEARING" in engine_service.MAPPING_LABELS, "clearing role visible in inventory account mappings")

# ===== ۴) تعریف فرایند =====
mods = ["GL", "INVENTORY", "RECEIVABLES", "PAYABLES", "CHEQUES", "FIXED_ASSETS"]
check(raises(lambda: orun.create_run(company_id, manager, "R1", "افتتاحیه", fy_id, balance_date, mods,
                                     go_live_date=fy_end + datetime.timedelta(days=5)), "همان سال مالی"),
      "go-live date must be inside the fiscal year")
check(raises(lambda: orun.create_run(company_id, manager, "R1", "افتتاحیه", fy_id, go_live, mods, go_live_date=go_live), "بعد از"),
      "balance date must be before go-live")
check(raises(lambda: orun.create_run(company_id, manager, "R1", "افتتاحیه", fy_id, balance_date, ["XX"]), "ماژول"),
      "unknown module rejected")
run = orun.create_run(company_id, manager, "R1", "افتتاحیهٔ ۱۴۰۴", fy_id, balance_date, mods)
r = orun.get_run(company_id, run)
check(r.status_code == "DRAFT" and r.go_live_date == go_live and len(r.modules) == 6 and r.progress == 0
      and all(m.status_code == "NOT_STARTED" for m in r.modules), "run created: draft, 6 modules not started, 0%")
check([m.module_code for m in r.modules] == ["INVENTORY", "RECEIVABLES", "PAYABLES", "CHEQUES", "FIXED_ASSETS", "GL"],
      "modules listed in posting order (GL last)")
check(raises(lambda: orun.create_run(company_id, manager, "R2", "دوباره", fy_id, balance_date, ["GL"]), "اصلاحی"),
      "second ORIGINAL run for the same fiscal year is refused")
check(raises(lambda: orun.create_run(company_id, manager, "R1", "تکرار کد", fy_id, balance_date, ["GL"]), "اصلاحی")
      or raises(lambda: orun.create_run(company_id, manager, "R1", "تکرار کد", fy_id, balance_date, ["GL"]), "تکراری"),
      "duplicate code refused")
check(raises(lambda: orun.create_run(company_id, uid, "C1", "اصلاحی", fy_id, balance_date, ["GL"], corrects_run_id=run), "ثبت‌شده"),
      "correction run only for a posted run")

check(raises(lambda: orun.set_people(company_id, manager, run, "GL", entry, entry), "متفاوت"),
      "owner and approver must differ (self-approval off)")
for m in mods:
    orun.set_people(company_id, manager, run, m, entry, approver)
check(all(m.owner_user_id == entry and m.approver_user_id == approver for m in orun.get_run(company_id, run).modules),
      "owner and approver assigned per module")

# ===== ۵) قالب =====
tpl = oim.write_template("GL", os.path.join(tmp, "gl_template.xlsx"))
import openpyxl  # noqa: E402
wb = openpyxl.load_workbook(tpl)
hdr = [c_.value for c_ in wb.worksheets[0][1]]
check(hdr[0] == "کد حساب معین *" and "بدهکار" in hdr and len(wb.sheetnames) == 2, f"xlsx template with required marks and guide sheet ({hdr[:3]})")
tpl_csv = oim.write_template("INVENTORY", os.path.join(tmp, "inv_template.csv"))
with open(tpl_csv, encoding="utf-8-sig") as fh:
    check(fh.readline().startswith("کد انبار *,کد کالا *,مقدار *,بهای واحد *"), "csv template header")

# ===== ۶) بارگذاری تراز با خطای ستونی =====
H = ["کد حساب معین", "بدهکار", "بستانکار", "کد تفصیلی شخص", "شرح", "ستون اضافه"]
gl_rows_v1 = [
    ("111", "1,500,000", "", "", "صندوق", "x"),
    ("311", "", "۲٬۰۰۰٬۰۰۰", "", "سرمایه", ""),
    ("131", "500000", "", "", "دریافتنی بدون شخص", ""),       # تفصیلی الزامی
    ("999", "100", "", "", "حساب ناموجود", ""),
    ("", "", "", "", "", ""),                                 # ردیف خالی ← نادیده
    ("12", "100", "", "", "حساب غیرمعین", ""),
    ("391", "100", "", "", "حساب واسط", ""),
    ("211", "50", "50", "S1", "هر دو ستون", ""),
]
f1 = write_xlsx(os.path.join(tmp, "gl_v1.xlsx"), H, gl_rows_v1)
check(raises(lambda: oim.import_file(company_id, viewer, run, "GL", f1), "دسترسی"), "viewer cannot import")
check(raises(lambda: oim.import_file(company_id, entry2, run, "GL", f1), "مسئول ورود"), "only the module owner enters data")
res = oim.import_file(company_id, entry, run, "GL", f1)
check(res.version_no == 1 and res.row_count == 7 and res.valid == 2 and res.invalid == 5,
      f"version 1: 7 data rows (blank skipped), 2 valid, 5 invalid ({res.row_count}/{res.valid}/{res.invalid})")
check(any("ستون اضافه" in n for n in res.notes), "unknown column reported, not fatal")
rows = {r.source_row_no: r for r in rows_of(run, "GL")}
err = lambda no: " | ".join(f"{e['label']}: {e['msg']}" for e in (rows[no].errors or []))  # noqa: E731
check("کد حساب معین" in err(4) and "شخص" in err(4), f"row 4: missing required person detail, reported on the account column ({err(4)})")
check("در کدینگ" in err(5), f"row 5: unknown account ({err(5)})")
check("قابل ثبت" in err(7), f"row 7: non-postable account ({err(7)})")
check("واسط" in err(8), f"row 8: clearing account not allowed in trial balance ({err(8)})")
check("نه هر دو" in err(9), f"row 9: debit and credit together ({err(9)})")
check(rows[3].amount == D(-2_000_000) and rows[2].amount == D(1_500_000), "Persian digits/separators parsed; signed base amount")
check(module_status(run, "GL") == "HAS_ERRORS" and orun.get_run(company_id, run).status_code == "IN_PROGRESS",
      "module HAS_ERRORS, run moved to IN_PROGRESS")

# همان فایل دوباره
check(raises(lambda: oim.import_file(company_id, entry, run, "GL", f1), "نسخهٔ 1"), "same file (same hash) cannot be uploaded twice")

# ===== ۷) بارگذاری اصلاحی: بدون ردیف تکراری =====
gl_rows_v2 = [
    ("111", "1500000", "", "", "صندوق", ""),
    ("311", "", "2000000", "", "سرمایه", ""),
    ("131", "500000", "", "C1", "دریافتنی مشتری یک", ""),
]
f2 = write_xlsx(os.path.join(tmp, "gl_v2.xlsx"), H, gl_rows_v2)
res2 = oim.import_file(company_id, entry, run, "GL", f2, replace_missing=True)
all_rows = rows_of(run, "GL")
check(res2.version_no == 2 and len(all_rows) == 3 and res2.removed == 5 and all(r.status_code == "VALID" for r in all_rows),
      f"version 2 upserts by key, removes the stale invalid rows; 3 valid rows, no duplicates ({len(all_rows)}, removed {res2.removed})")
check(module_status(run, "GL") == "ENTERING", "module back to ENTERING when no errors remain")
with new_session() as s:
    b = s.scalars(select(ImportBatch).where(ImportBatch.run_id == run, ImportBatch.module_code == "GL").order_by(ImportBatch.version_no)).all()
check(len(b) == 2 and b[0].file_name == "gl_v1.xlsx" and len(b[0].file_hash) == 64 and b[1].uploaded_by_user_id == entry
      and b[0].invalid_count == 5, "each upload recorded: version, file name, SHA-256, user, counts")

# تکرار کلید در یک فایل
f3 = write_csv(os.path.join(tmp, "gl_dup.csv"), H, [("111", "10", "", "", "a", ""), ("111", "20", "", "", "b", "")])
res3 = oim.import_file(company_id, entry, run, "GL", f3)
dup_row = next(r for r in rows_of(run, "GL") if r.row_key.startswith("1-11-111|"))
check(res3.duplicates == [(2, 3)] and dup_row.status_code == "INVALID" and "ردیف 3" in str(dup_row.errors),
      "duplicate key inside one file: flagged on the first row, second not stored")
check(len(rows_of(run, "GL")) == 3, "still 3 rows (no duplicate created)")

# ===== ۸) ورود دستی =====
check(raises(lambda: oim.save_row(company_id, entry, run, "GL", {"account_code": "311", "credit": "5"}), "کلید"),
      "form entry with an existing key is refused (edit instead)")
fixed = oim.save_row(company_id, entry, run, "GL", {"account_code": "111", "debit": "1500000", "description": "صندوق"},
                     row_id=dup_row.row_id)
check(fixed.status == "VALID" and module_status(run, "GL") == "ENTERING", "editing the flagged row fixes it")
extra = oim.save_row(company_id, entry, run, "GL", {"account_code": "211", "credit": "0", "person_code": "S1"})
check(extra.status == "INVALID" and module_status(run, "GL") == "HAS_ERRORS", "invalid form row saved for later fix")
oim.delete_row(company_id, entry, run, "GL", extra.row_id)
check(module_status(run, "GL") == "ENTERING", "deleting it clears the error state")

# ===== ۹) اعتبارسنجی ماژول‌های دیگر =====
dims.set_account_dimension_types(cash_gl.account_id, company_id, [cc_type])
rv = oim.revalidate(company_id, entry, run, "GL")
check(rv.invalid == 1 and "مرکز هزینه" in str(next(r for r in rows_of(run, "GL") if r.row_key.startswith("1-11-111|")).errors),
      "revalidate after master-data change: cash now needs a cost center")
dims.set_account_dimension_types(cash_gl.account_id, company_id, [])
oim.revalidate(company_id, entry, run, "GL")

batch_item = item("BT-1", "کالای بچ‌دار", "GOOD", track_batch=True, track_expiry=True)
IH = ["کد انبار", "کد کالا", "مقدار", "بهای واحد", "شمارهٔ بچ", "تاریخ انقضا"]
f_inv = write_csv(os.path.join(tmp, "inv.csv"), IH, [
    ("RM", "RM-1", "10", "", "", ""),            # بهای خالی ← خطا (بها ساخته نمی‌شود)
    ("RM", "RM-2", "5", "0", "", ""),            # بهای صفر ← هشدار
    ("RM", "BT-1", "3", "1000", "", ""),         # بچ و انقضا الزامی
    ("RM", "BT-1", "4", "1000", "B-7", "1406/01/01"),
    ("ZZ", "RM-3", "1", "1", "", ""),
    ("RM", "RM-3", "-1", "1", "", ""),
])
ri = oim.import_file(company_id, entry, run, "INVENTORY", f_inv)
inv = {r.source_row_no: r for r in rows_of(run, "INVENTORY")}
check("بهای واحد" in str(inv[2].errors) and inv[2].status_code == "INVALID", "missing unit cost is an error, never a fabricated cost")
check(inv[3].status_code == "VALID" and "بهای صفر" in str(inv[3].warnings), "zero cost accepted with a warning")
check("بچ" in str(inv[4].errors) and "انقضا" in str(inv[4].errors), "batch/expiry required for tracked item")
check(inv[5].status_code == "VALID" and inv[5].resolved["batch_no"] == "B-7" and inv[5].amount == D(4000)
      and inv[5].resolved["expiry_date"].startswith("2027"), "batch row valid; Jalali expiry parsed; amount = qty × cost")
check("انبار" in str(inv[6].errors) and "صفر" in str(inv[7].errors), "unknown warehouse / non-positive quantity")
check(ri.invalid == 4 and module_status(run, "INVENTORY") == "HAS_ERRORS", "inventory module has errors")

ARH = ["کد تفصیلی مشتری", "شمارهٔ سند در سیستم قبلی", "تاریخ سند", "سررسید", "مبلغ اصلی", "تسویه‌شده تا برش", "ماندهٔ باز"]
bd = fmt_date = oc.fmt_date
f_ar = write_csv(os.path.join(tmp, "ar.csv"), ARH, [
    ("C1", "F-100", bd(balance_date - datetime.timedelta(days=40)), bd(balance_date + datetime.timedelta(days=20)), "800000", "300000", "500000"),
    ("C2", "F-200", bd(balance_date), "", "100000", "100000", ""),      # کاملاً تسویه‌شده
    ("C2", "F-201", bd(balance_date), "", "100000", "", "90000"),       # ماندهٔ ناسازگار
    ("C2", "F-202", bd(go_live), "", "100000", "", ""),                 # بعد از برش
    ("C9", "F-300", bd(balance_date), "", "100000", "", ""),
])
oim.import_file(company_id, entry, run, "RECEIVABLES", f_ar)
ar = {r.source_row_no: r for r in rows_of(run, "RECEIVABLES")}
check(ar[2].status_code == "VALID" and ar[2].amount == D(500000) and ar[2].resolved["settled"] == "300000",
      "open item: original 800,000 − settled 300,000 = open 500,000")
check("تسویه شده" in str(ar[3].errors), "fully settled document is not transferred")
check("برابر نیست" in str(ar[4].errors), "open amount must equal original − settled")
check("بعد از تاریخ مانده" in str(ar[5].errors), "document after the balance date refused")
check("پیدا نشد" in str(ar[6].errors), "unknown counterparty")

f_chq = write_csv(os.path.join(tmp, "chq.csv"), ["نوع چک", "شمارهٔ چک", "سررسید", "مبلغ", "کد تفصیلی طرف‌حساب"], [
    ("دریافتی", "123456", bd(go_live + datetime.timedelta(days=30)), "250000", "C1"),
    ("نامعلوم", "1", bd(go_live), "1", "C1"),
])
oim.import_file(company_id, entry, run, "CHEQUES", f_chq)
chq = {r.source_row_no: r for r in rows_of(run, "CHEQUES")}
check("نگاشت نشده" in str(chq[2].errors), "cheque refused while the treasury cheque account is not mapped")
check("مجاز نیست" in str(chq[3].errors), "invalid choice value reported with allowed values")

from peecha.services.fixed_assets import common as fac  # noqa: E402
cat_ok = fac.save_category(company_id, fac.CategoryFields("MCH", "ماشین‌آلات", default_life_months=120, accounts={
    "asset_account_id": fa_cost_gl.account_id, "accumulated_depreciation_account_id": fa_dep_gl.account_id,
    "depreciation_expense_account_id": dep_exp_gl.account_id}))
fac.save_category(company_id, fac.CategoryFields("BAD", "بدون حساب"))
FH = ["کد دارایی", "نام دارایی", "کد طبقه", "بهای تمام‌شده", "استهلاک انباشته تا برش", "تاریخ شروع استهلاک/بهره‌برداری"]
start = fac.add_months(balance_date, -23)
f_fa = write_csv(os.path.join(tmp, "fa.csv"), FH, [
    ("M-1", "دستگاه برش", "MCH", "1200000", "240000", bd(start)),
    ("M-2", "بدون حساب", "BAD", "100", "", bd(start)),
    ("M-3", "استهلاک زیاد", "MCH", "100", "200", bd(start)),
])
oim.import_file(company_id, entry, run, "FIXED_ASSETS", f_fa)
fa = {r.source_row_no: r for r in rows_of(run, "FIXED_ASSETS")}
check(fa[2].status_code == "VALID" and fa[2].amount == D(960000) and fa[2].resolved["elapsed_months"] == 24
      and fa[2].resolved["life_months"] == "120", "asset: book value 960,000, 24 months elapsed to cutover, life from category")
check("حساب" in str(fa[3].errors), "category without accounts refused")
check("بیشتر است" in str(fa[4].errors), "accumulated depreciation above cost refused")

# ===== ۱۰) گردش وضعیت و تفکیک وظایف =====
check(raises(lambda: orun.submit(company_id, entry, run, "INVENTORY"), "خطا"), "module with errors cannot be submitted")
check(raises(lambda: orun.submit(company_id, entry2, run, "GL"), "مسئول ورود"), "only the owner submits")
orun.submit(company_id, entry, run, "GL", "تراز کامل است")
check(module_status(run, "GL") == "READY", "GL ready for review")
check(raises(lambda: orun.approve(company_id, entry, run, "GL"), "دسترسی"), "entry role lacks approval permission")
toggle(oc.OPN_SELF_APPROVAL, True)
orun.set_people(company_id, manager, run, "PAYABLES", approver, approver)
toggle(oc.OPN_SELF_APPROVAL, False)
check(raises(lambda: orun.reject(company_id, approver, run, "GL", " "), "دلیل"), "rejection needs a reason")
orun.reject(company_id, approver, run, "GL", "حساب صندوق را بررسی کنید")
gm = next(m for m in orun.get_run(company_id, run).modules if m.module_code == "GL")
check(gm.status_code == "ENTERING" and gm.reject_reason == "حساب صندوق را بررسی کنید", "rejected back to entering with reason")
orun.submit(company_id, entry, run, "GL")
orun.approve(company_id, approver, run, "GL", "تأیید شد")
gm = next(m for m in orun.get_run(company_id, run).modules if m.module_code == "GL")
check(gm.status_code == "APPROVED" and gm.approved_by_user_id == approver and gm.submitted_by_user_id == entry,
      "approved by a different user than the submitter")
oim.save_row(company_id, entry, run, "GL", {"account_code": "121", "debit": "1000"})
with new_session() as s:
    resets = s.scalar(select(func.count()).select_from(OpeningEvent).where(OpeningEvent.run_id == run, OpeningEvent.event_code == "RESET"))
check(module_status(run, "GL") == "ENTERING" and resets == 1, "data change after approval returns the module to entering (RESET event)")
check(next(r for r in rows_of(run, "GL") if r.row_key.startswith("1-12-121|")).warnings, "control account (inventory) flagged with a warning")

# تفکیک: ارسال‌کننده نمی‌تواند تأیید کند حتی با نقش تأیید
both = user("opn_both", "OPN_ENTRY", "OPN_APPROVER")
orun.set_people(company_id, manager, run, "GL", both, None)
orun.submit(company_id, both, run, "GL")
check(raises(lambda: orun.approve(company_id, both, run, "GL"), "تفکیک وظایف"), "submitter cannot approve own data")
toggle(oc.OPN_SELF_APPROVAL, True)
orun.approve(company_id, both, run, "GL")
toggle(oc.OPN_SELF_APPROVAL, False)
check(module_status(run, "GL") == "APPROVED", "self-approval only when the toggle is on")
check(orun.get_run(company_id, run).progress > 0, "run progress reflects module states")

# ===== ۱۱) حفاظت ردیف ثبت‌شده (تریگر پایگاه داده) =====
posted = next(r for r in rows_of(run, "FIXED_ASSETS") if r.source_row_no == 2)
with new_session() as s:
    s.execute(text("UPDATE opn.staged_rows SET status_code = 'POSTED' WHERE row_id = :r"), {"r": posted.row_id})
    s.commit()


def sql(stmt):
    with new_session() as s:
        s.execute(text(stmt), {"r": posted.row_id})
        s.commit()


try:
    sql("UPDATE opn.staged_rows SET payload = '{}'::jsonb WHERE row_id = :r")
    blocked_update = False
except DBAPIError:
    blocked_update = True
try:
    sql("DELETE FROM opn.staged_rows WHERE row_id = :r")
    blocked_delete = False
except DBAPIError:
    blocked_delete = True
check(blocked_update and blocked_delete, "database trigger blocks editing or deleting a posted row")
check(raises(lambda: oim.delete_row(company_id, entry, run, "FIXED_ASSETS", posted.row_id), "ثبت‌شده"), "service refuses too")
f_fa2 = write_csv(os.path.join(tmp, "fa2.csv"), FH, [("M-1", "دستگاه برش (ویرایش)", "MCH", "999", "", bd(start))])
rfa = oim.import_file(company_id, entry, run, "FIXED_ASSETS", f_fa2)
after = next(r for r in rows_of(run, "FIXED_ASSETS") if r.row_id == posted.row_id)
check(rfa.posted_skipped == [2] and after.payload["cost"] == "1200000", "re-upload never overwrites a posted row")
check(raises(lambda: orun.set_modules(company_id, manager, run, ["GL"]), "اطلاعات واردشده"),
      "a module with data cannot be removed from the run")
check(raises(lambda: orun.cancel_run(company_id, manager, run, "تست"), "اصلاحی"), "run with posted rows cannot be cancelled")
with new_session() as s:
    s.execute(text("ALTER TABLE opn.staged_rows DISABLE TRIGGER tr_opn_rows_guard"))
    s.execute(text("UPDATE opn.staged_rows SET status_code = 'VALID' WHERE row_id = :r"), {"r": posted.row_id})
    s.execute(text("ALTER TABLE opn.staged_rows ENABLE TRIGGER tr_opn_rows_guard"))
    s.commit()

# ===== ۱۲) بارگذاری هم‌زمان =====
fa_a = write_csv(os.path.join(tmp, "fa_a.csv"), FH, [("M-10", "الف", "MCH", "100", "", bd(start))])
fa_b = write_csv(os.path.join(tmp, "fa_b.csv"), FH, [("M-10", "ب", "MCH", "200", "", bd(start))])
errs = []


def up(path):
    try:
        oim.import_file(company_id, entry, run, "FIXED_ASSETS", path)
    except Exception as exc:  # noqa: BLE001
        errs.append(exc)


ts = [threading.Thread(target=up, args=(p,)) for p in (fa_a, fa_b)]
[t.start() for t in ts]
[t.join() for t in ts]
with new_session() as s:
    versions = sorted(s.scalars(select(ImportBatch.version_no).where(ImportBatch.run_id == run, ImportBatch.module_code == "FIXED_ASSETS")))
m10 = [r for r in rows_of(run, "FIXED_ASSETS") if r.row_key == "M-10"]
check(not errs and len(m10) == 1 and versions == [1, 2, 3, 4], f"two concurrent uploads serialize: one row per key, versions {versions}")

# ===== ۱۳) لغو و فرایند تازه =====
with new_session() as s:
    audit_imports = s.scalar(select(func.count()).select_from(ActivityLog).where(
        ActivityLog.entity_type == oc.ENTITY_RUN, ActivityLog.entity_id == run))
check(audit_imports >= 8, f"run operations audited ({audit_imports})")
check(raises(lambda: orun.cancel_run(company_id, manager, run, ""), "دلیل"), "cancel needs a reason")
orun.cancel_run(company_id, manager, run, "شروع دوباره با کدینگ نهایی")
check(orun.get_run(company_id, run).status_code == "CANCELLED", "run cancelled")
check(raises(lambda: oim.save_row(company_id, entry, run, "GL", {"account_code": "111", "debit": "1"}), "پیش از ثبت"),
      "cancelled run accepts no data")
run2 = orun.create_run(company_id, manager, "R2", "افتتاحیهٔ نهایی", fy_id, balance_date, ["GL"])
check(run2 > run and [x.code for x in orun.list_runs(company_id)][:2] == ["R2", "R1"],
      "after cancelling, a new ORIGINAL run for the same fiscal year is allowed")
with new_session() as s:
    m = s.get(RunModule, (run2, "GL"))
check(m is not None and m.status_code == "NOT_STARTED", "new run starts clean")

fx.finish()
